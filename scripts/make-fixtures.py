"""Teknik denemedeki Office → PDF testleri için örnek Word, Excel ve PowerPoint dosyaları üretir.

Gerekenler: python-docx, openpyxl, python-pptx. Çıktılar Fixtures/ altına yazılır ve depoya alınır.
"""
from pathlib import Path

import docx
import openpyxl
import pptx

OUT = Path(__file__).resolve().parent.parent / "Fixtures"
TURKISH = "Türkçe karakterler: ğüşıöç İĞÜŞÖÇ"


def main():
    OUT.mkdir(exist_ok=True)

    document = docx.Document()
    document.add_heading("PDF Atölye Word deneme", level=1)
    document.add_paragraph(TURKISH)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Araç", "Durum"), ("Birleştir", "Hazır"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(OUT / "ornek.docx")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Deneme"
    sheet["A1"] = "PDF Atölye Excel deneme"
    sheet.append(["Araç", "Adet"])
    sheet.append(["Birleştir", 12])
    sheet.append(["Sıkıştır", 7])
    sheet["A5"] = TURKISH
    workbook.save(OUT / "ornek.xlsx")

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "PDF Atölye PowerPoint deneme"
    slide.placeholders[1].text = TURKISH
    deck.save(OUT / "ornek.pptx")


if __name__ == "__main__":
    main()
